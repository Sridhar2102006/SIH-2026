import openpyxl, glob, os, json

result = {}
files = glob.glob('*.xlsx')
for f in files:
    if f.startswith('~$'): continue
    print(f"Reading {f}...")
    wb = openpyxl.load_workbook(f, read_only=True)
    file_info = {'sheets': wb.sheetnames, 'sample_data': {}}
    for name in wb.sheetnames:
        ws = wb[name]
        rows = []
        for i, row in enumerate(ws.iter_rows(values_only=True)):
            if i > 8: break
            row_clean = [str(c) if c is not None else None for c in row]
            rows.append(row_clean)
        file_info['sample_data'][name] = rows
    result[f] = file_info

with open('scripts/excel_summary.json', 'w', encoding='utf-8') as out:
    json.dump(result, out, indent=2, ensure_ascii=False)

print("Saved scripts/excel_summary.json successfully.")
