import openpyxl, json, os, zlib, struct, xml.etree.ElementTree as ET

os.makedirs('src/data/matrices', exist_ok=True)

# 1. 20 Curated Workspace Visual Identity Families (Section 21-23 of specification)
THEME_PALETTES = [
    {
        "themeId": "theme-01-honey-gold",
        "name": "Honey Gold",
        "family": "Warm Apiary",
        "primaryColor": "#D99A24",
        "secondaryColor": "#B87316",
        "accentColor": "#F3B749",
        "surfaceColor": "#FFFDF8",
        "backgroundColor": "#FFF9EF",
        "textColor": "#34261B",
        "iconAccent": "#D99A24",
        "navigationAccent": "#D99A24",
        "cardBorder": "#EDE2D1",
        "headerTreatment": "warm-gold-gradient"
    },
    {
        "themeId": "theme-02-sage",
        "name": "Sage",
        "family": "Herbal Colony",
        "primaryColor": "#4F7A52",
        "secondaryColor": "#3E6340",
        "accentColor": "#71845B",
        "surfaceColor": "#FBFDFB",
        "backgroundColor": "#F4F7F2",
        "textColor": "#233324",
        "iconAccent": "#4F7A52",
        "navigationAccent": "#4F7A52",
        "cardBorder": "#D6E3D4",
        "headerTreatment": "sage-leaf-gradient"
    },
    {
        "themeId": "theme-03-forest",
        "name": "Forest",
        "family": "Deep Woodland",
        "primaryColor": "#2C5E3B",
        "secondaryColor": "#1E4229",
        "accentColor": "#49855C",
        "surfaceColor": "#FAFCFA",
        "backgroundColor": "#F0F5F1",
        "textColor": "#1A2E20",
        "iconAccent": "#2C5E3B",
        "navigationAccent": "#2C5E3B",
        "cardBorder": "#CEE0D2",
        "headerTreatment": "deep-forest-gradient"
    },
    {
        "themeId": "theme-04-olive",
        "name": "Olive",
        "family": "Mediterranean Orchard",
        "primaryColor": "#6E7A36",
        "secondaryColor": "#545E28",
        "accentColor": "#91A049",
        "surfaceColor": "#FCFDFC",
        "backgroundColor": "#F7F8F0",
        "textColor": "#2B3016",
        "iconAccent": "#6E7A36",
        "navigationAccent": "#6E7A36",
        "cardBorder": "#E2E7CC",
        "headerTreatment": "olive-grove-gradient"
    },
    {
        "themeId": "theme-05-terracotta",
        "name": "Terracotta",
        "family": "Earthen Clay",
        "primaryColor": "#B85450",
        "secondaryColor": "#943F3C",
        "accentColor": "#D9736F",
        "surfaceColor": "#FFFDFD",
        "backgroundColor": "#FDF6F5",
        "textColor": "#3B1E1D",
        "iconAccent": "#B85450",
        "navigationAccent": "#B85450",
        "cardBorder": "#EED8D6",
        "headerTreatment": "terracotta-sun-gradient"
    },
    {
        "themeId": "theme-06-amber",
        "name": "Amber",
        "family": "Raw Honeycomb",
        "primaryColor": "#C97A1E",
        "secondaryColor": "#A15E12",
        "accentColor": "#E89B3A",
        "surfaceColor": "#FFFDF9",
        "backgroundColor": "#FCF6ED",
        "textColor": "#3B2610",
        "iconAccent": "#C97A1E",
        "navigationAccent": "#C97A1E",
        "cardBorder": "#EFE0CA",
        "headerTreatment": "amber-crystal-gradient"
    },
    {
        "themeId": "theme-07-copper",
        "name": "Copper",
        "family": "Artisan Settling",
        "primaryColor": "#A85D32",
        "secondaryColor": "#854521",
        "accentColor": "#CA7A4C",
        "surfaceColor": "#FFFDFB",
        "backgroundColor": "#FAF3EE",
        "textColor": "#361F13",
        "iconAccent": "#A85D32",
        "navigationAccent": "#A85D32",
        "cardBorder": "#EADCD2",
        "headerTreatment": "burnished-copper-gradient"
    },
    {
        "themeId": "theme-08-sand",
        "name": "Sand",
        "family": "Desert Flora",
        "primaryColor": "#997D54",
        "secondaryColor": "#7A623F",
        "accentColor": "#BC9F73",
        "surfaceColor": "#FFFEFC",
        "backgroundColor": "#F9F6F0",
        "textColor": "#362B1C",
        "iconAccent": "#997D54",
        "navigationAccent": "#997D54",
        "cardBorder": "#E8DFD0",
        "headerTreatment": "sandstone-gradient"
    },
    {
        "themeId": "theme-09-teal",
        "name": "Teal",
        "family": "Coastal Apiary",
        "primaryColor": "#2E7C7F",
        "secondaryColor": "#215E61",
        "accentColor": "#4AA1A4",
        "surfaceColor": "#FBFCFD",
        "backgroundColor": "#F1F7F8",
        "textColor": "#142F30",
        "iconAccent": "#2E7C7F",
        "navigationAccent": "#2E7C7F",
        "cardBorder": "#CEE3E4",
        "headerTreatment": "coastal-teal-gradient"
    },
    {
        "themeId": "theme-10-deep-blue",
        "name": "Deep Blue",
        "family": "Highland Purity",
        "primaryColor": "#24587D",
        "secondaryColor": "#1A415E",
        "accentColor": "#3E7CA9",
        "surfaceColor": "#FAFCFE",
        "backgroundColor": "#F0F5FA",
        "textColor": "#122638",
        "iconAccent": "#24587D",
        "navigationAccent": "#24587D",
        "cardBorder": "#CDDFED",
        "headerTreatment": "highland-sky-gradient"
    },
    {
        "themeId": "theme-11-indigo",
        "name": "Indigo",
        "family": "Evening Flight",
        "primaryColor": "#434B8C",
        "secondaryColor": "#32396E",
        "accentColor": "#616AC0",
        "surfaceColor": "#FBFBFE",
        "backgroundColor": "#F3F4FB",
        "textColor": "#1B1E3B",
        "iconAccent": "#434B8C",
        "navigationAccent": "#434B8C",
        "cardBorder": "#D6D9F0",
        "headerTreatment": "indigo-dusk-gradient"
    },
    {
        "themeId": "theme-12-plum",
        "name": "Plum",
        "family": "Wild Berry Blossom",
        "primaryColor": "#7B447D",
        "secondaryColor": "#603262",
        "accentColor": "#A15EA3",
        "surfaceColor": "#FDFBFE",
        "backgroundColor": "#F8F2F8",
        "textColor": "#331834",
        "iconAccent": "#7B447D",
        "navigationAccent": "#7B447D",
        "cardBorder": "#E9D7EA",
        "headerTreatment": "wild-plum-gradient"
    },
    {
        "themeId": "theme-13-berry",
        "name": "Berry",
        "family": "Nectar Bloom",
        "primaryColor": "#993B63",
        "secondaryColor": "#792C4D",
        "accentColor": "#BE5783",
        "surfaceColor": "#FDFBFC",
        "backgroundColor": "#FAF1F5",
        "textColor": "#3B1424",
        "iconAccent": "#993B63",
        "navigationAccent": "#993B63",
        "cardBorder": "#EED3DF",
        "headerTreatment": "berry-blossom-gradient"
    },
    {
        "themeId": "theme-14-coral",
        "name": "Coral",
        "family": "Warm Pollen",
        "primaryColor": "#B85843",
        "secondaryColor": "#944230",
        "accentColor": "#D97660",
        "surfaceColor": "#FFFEFD",
        "backgroundColor": "#FCF4F1",
        "textColor": "#3D1A13",
        "iconAccent": "#B85843",
        "navigationAccent": "#B85843",
        "cardBorder": "#EFDACF",
        "headerTreatment": "pollen-coral-gradient"
    },
    {
        "themeId": "theme-15-slate",
        "name": "Slate",
        "family": "Laboratory Stainless",
        "primaryColor": "#4F6370",
        "secondaryColor": "#3B4B54",
        "accentColor": "#6E8799",
        "surfaceColor": "#FCFDFE",
        "backgroundColor": "#F2F5F7",
        "textColor": "#1C252B",
        "iconAccent": "#4F6370",
        "navigationAccent": "#4F6370",
        "cardBorder": "#D1DDE4",
        "headerTreatment": "stainless-slate-gradient"
    },
    {
        "themeId": "theme-16-ocean",
        "name": "Ocean",
        "family": "Clean Transport",
        "primaryColor": "#1E6C8E",
        "secondaryColor": "#155069",
        "accentColor": "#328EBA",
        "surfaceColor": "#FBFCFD",
        "backgroundColor": "#EFF6F9",
        "textColor": "#0E2936",
        "iconAccent": "#1E6C8E",
        "navigationAccent": "#1E6C8E",
        "cardBorder": "#CCE2ED",
        "headerTreatment": "ocean-transit-gradient"
    },
    {
        "themeId": "theme-17-moss",
        "name": "Moss",
        "family": "Shaded Hive Yard",
        "primaryColor": "#436E46",
        "secondaryColor": "#325435",
        "accentColor": "#5E9662",
        "surfaceColor": "#FAFDFB",
        "backgroundColor": "#F1F7F2",
        "textColor": "#172E19",
        "iconAccent": "#436E46",
        "navigationAccent": "#436E46",
        "cardBorder": "#CEE4D1",
        "headerTreatment": "shaded-moss-gradient"
    },
    {
        "themeId": "theme-18-clay",
        "name": "Clay",
        "family": "Ceramic Pot",
        "primaryColor": "#96593B",
        "secondaryColor": "#77432A",
        "accentColor": "#B77553",
        "surfaceColor": "#FFFDFB",
        "backgroundColor": "#F9F4F0",
        "textColor": "#361D12",
        "iconAccent": "#96593B",
        "navigationAccent": "#96593B",
        "cardBorder": "#E7D8CF",
        "headerTreatment": "natural-clay-gradient"
    },
    {
        "themeId": "theme-19-warm-brown",
        "name": "Warm Brown",
        "family": "Cedar Woodware",
        "primaryColor": "#6E4E37",
        "secondaryColor": "#543A27",
        "accentColor": "#8F674A",
        "surfaceColor": "#FFFDFC",
        "backgroundColor": "#F8F3EF",
        "textColor": "#2B1D14",
        "iconAccent": "#6E4E37",
        "navigationAccent": "#6E4E37",
        "cardBorder": "#E4D6CC",
        "headerTreatment": "cedar-box-gradient"
    },
    {
        "themeId": "theme-20-charcoal",
        "name": "Charcoal",
        "family": "Carbon Ledger Trust",
        "primaryColor": "#3D4147",
        "secondaryColor": "#2A2D31",
        "accentColor": "#5B616B",
        "surfaceColor": "#FDFDFD",
        "backgroundColor": "#F4F5F6",
        "textColor": "#191B1D",
        "iconAccent": "#3D4147",
        "navigationAccent": "#3D4147",
        "cardBorder": "#D6D9DE",
        "headerTreatment": "carbon-trust-gradient"
    }
]

with open('src/data/matrices/theme_palettes.json', 'w', encoding='utf-8') as f:
    json.dump(THEME_PALETTES, f, indent=2)

print("Saved 20 curated themes to src/data/matrices/theme_palettes.json")

# 2. Extract Beekeeper Capabilities, Policies, Onboarding
wb_bee = openpyxl.load_workbook('HoneyChain_Beekeeper_Feature_Matrix.xlsx', read_only=True)
beekeeper_caps = []
for r in list(wb_bee['Capability Catalog'].iter_rows(values_only=True))[1:]:
    if r[0]:
        beekeeper_caps.append({
            "id": r[0],
            "designationFamily": "BEEKEEPER",
            "name": r[1],
            "type": r[2],
            "purpose": r[3],
            "category": "FIELD"
        })

beekeeper_onboarding = []
for r in list(wb_bee['Onboarding Mapping'].iter_rows(values_only=True))[1:]:
    if r[0]:
        beekeeper_onboarding.append({
            "designationFamily": "BEEKEEPER",
            "question": r[0],
            "answerSignal": r[1],
            "capabilityTrigger": r[2],
            "resultingFeatures": r[3]
        })

beekeeper_policies = []
for r in list(wb_bee['Beekeeper Policy'].iter_rows(values_only=True))[1:]:
    if r[0]:
        beekeeper_policies.append({
            "designationFamily": "BEEKEEPER",
            "policy": r[0],
            "rule": r[1]
        })

beekeeper_features = []
seen_bee_features = set()
for r in list(wb_bee['Feature Matrix'].iter_rows(values_only=True))[1:]:
    if r[7] == 'YES' and r[4] and r[6]:
        key = (r[4], r[6])
        if key not in seen_bee_features:
            seen_bee_features.add(key)
            beekeeper_features.append({
                "designationFamily": "BEEKEEPER",
                "featureType": r[2] or "Feature",
                "name": r[3],
                "requiredCapability": r[4],
                "accessLevel": r[5] or "Standard",
                "permissionId": r[6]
            })

# 3. Extract Processor Capabilities, Features
with open('HoneyChain_Processor_Feature_Matrix.xlsx', 'rb') as f:
    proc_raw = f.read()

def get_proc_sheet(num_str):
    pos = 0
    while True:
        idx = proc_raw.find(b'PK\x03\x04', pos)
        if idx == -1: break
        sig, ver, flags, method, mtime, mdate, crc, comp_size, uncomp_size, fn_len, extra_len = struct.unpack('<IHHHHHIIIHH', proc_raw[idx:idx+30])
        fn = proc_raw[idx+30:idx+30+fn_len].decode('latin1')
        data_start = idx + 30 + fn_len + extra_len
        if num_str in fn:
            raw = proc_raw[data_start:data_start+comp_size]
            return zlib.decompress(raw, -15).decode('utf-8', 'replace')
        pos = data_start + comp_size
    return None

sheet2_xml = get_proc_sheet('sheet2.xml')
sheet3_xml = get_proc_sheet('sheet3.xml')

processor_caps = []
if sheet2_xml:
    root = ET.fromstring(sheet2_xml)
    for i, row in enumerate(root.iter('{http://schemas.openxmlformats.org/spreadsheetml/2006/main}row')):
        if i == 0: continue
        txts = [c.find('.//{http://schemas.openxmlformats.org/spreadsheetml/2006/main}t') for c in row.iter('{http://schemas.openxmlformats.org/spreadsheetml/2006/main}c')]
        vals = [t.text if t is not None else '' for t in txts]
        if vals and vals[0]:
            processor_caps.append({
                "id": vals[0],
                "designationFamily": "PROCESSOR",
                "name": vals[1],
                "type": vals[2],
                "purpose": vals[3],
                "category": "PRODUCTION"
            })

processor_features = []
if sheet3_xml:
    root = ET.fromstring(sheet3_xml)
    for i, row in enumerate(root.iter('{http://schemas.openxmlformats.org/spreadsheetml/2006/main}row')):
        if i == 0: continue
        txts = [c.find('.//{http://schemas.openxmlformats.org/spreadsheetml/2006/main}t') for c in row.iter('{http://schemas.openxmlformats.org/spreadsheetml/2006/main}c')]
        vals = [t.text if t is not None else '' for t in txts]
        if vals and len(vals) >= 5 and vals[0]:
            processor_features.append({
                "designationFamily": "PROCESSOR",
                "featureType": vals[0],
                "name": vals[1],
                "requiredCapability": vals[2],
                "accessLevel": vals[3],
                "permissionId": vals[4]
            })

processor_onboarding = [
    {
        "designationFamily": "PROCESSOR",
        "question": "What do you work with?",
        "answerSignal": "Harvested honey / extraction facility",
        "capabilityTrigger": "PROCESSING_MANAGEMENT",
        "resultingFeatures": "Processing Workspace, Start/Edit Processing, Processing History"
    },
    {
        "designationFamily": "PROCESSOR",
        "question": "Do you intake harvested honey batches?",
        "answerSignal": "Yes",
        "capabilityTrigger": "BATCH_INTAKE",
        "resultingFeatures": "Batch Intake Queue, Accept Batch for Processing"
    },
    {
        "designationFamily": "PROCESSOR",
        "question": "Do you record processing steps (filtration, settling, temperature)?",
        "answerSignal": "Yes",
        "capabilityTrigger": "PROCESSING_STEP_RECORD",
        "resultingFeatures": "Record Processing Step, Edit Processing Step"
    },
    {
        "designationFamily": "PROCESSOR",
        "question": "Do you record extraction parameters (speed, moisture, tank pressure)?",
        "answerSignal": "Yes",
        "capabilityTrigger": "PROCESSING_PARAMETERS",
        "resultingFeatures": "Processing Parameters Recording"
    },
    {
        "designationFamily": "PROCESSOR",
        "question": "Do you capture processing evidence photos?",
        "answerSignal": "Yes",
        "capabilityTrigger": "PROCESSING_EVIDENCE",
        "resultingFeatures": "Processing Evidence Capture, Evidence Gallery"
    },
    {
        "designationFamily": "PROCESSOR",
        "question": "Do you trace batches back to harvest collections?",
        "answerSignal": "Yes",
        "capabilityTrigger": "BATCH_TRACEABILITY",
        "resultingFeatures": "Batch Traceability, Source Collection Details"
    },
    {
        "designationFamily": "PROCESSOR",
        "question": "Do you split or combine extraction batches?",
        "answerSignal": "Yes",
        "capabilityTrigger": "BATCH_SPLIT_MERGE",
        "resultingFeatures": "Split Batch, Merge Batches"
    },
    {
        "designationFamily": "PROCESSOR",
        "question": "Do you prepare batches for Quality lab testing?",
        "answerSignal": "Yes",
        "capabilityTrigger": "QUALITY_HANDOFF",
        "resultingFeatures": "Prepare Quality Handoff, Quality Handoff Status"
    },
    {
        "designationFamily": "PROCESSOR",
        "question": "Do you prepare completed batches for Bottling / Packaging?",
        "answerSignal": "Yes",
        "capabilityTrigger": "PACKAGING_HANDOFF",
        "resultingFeatures": "Prepare Packaging Handoff, Packaging Handoff Status"
    }
]

processor_policies = [
    {"designationFamily": "PROCESSOR", "policy": "Designation", "rule": "Processor is a facility production identity; it does not directly grant Hive management or Quality decisions."},
    {"designationFamily": "PROCESSOR", "policy": "Authorization", "rule": "Permissions are derived from confirmed processing capabilities + policy + workspace scope."},
    {"designationFamily": "PROCESSOR", "policy": "Base focus", "rule": "Batch intake, centrifugal extraction, settling tank logs, evidence capture and batch advancement."},
    {"designationFamily": "PROCESSOR", "policy": "Optional split/merge", "rule": "Batch Split / Merge is enabled only when multi-lot combining is confirmed."},
    {"designationFamily": "PROCESSOR", "policy": "Quality boundary", "rule": "Processor prepares Quality Handoffs but cannot certify or make quality pass/fail decisions."},
    {"designationFamily": "PROCESSOR", "policy": "Packaging boundary", "rule": "Processor prepares Packaging Handoffs but does not generate consumer QR identities without packaging authority."}
]

# 4. Extract Lab Capabilities, Features, Variants, Policies
wb_lab1 = openpyxl.load_workbook('HoneyChain_Lab_Designation_Feature_Matrix_PART_1.xlsx', read_only=True)
wb_lab2 = openpyxl.load_workbook('HoneyChain_Lab_Designation_Feature_Matrix_PART_2.xlsx', read_only=True)

lab_caps = []
for r in list(wb_lab1['Capability Catalog'].iter_rows(values_only=True))[1:]:
    if r[0]:
        lab_caps.append({
            "id": r[0],
            "designationFamily": "LAB",
            "name": r[1],
            "type": r[2],
            "purpose": r[3],
            "category": "QUALITY"
        })

for r in list(wb_lab2['Advanced Capability Catalog'].iter_rows(values_only=True))[1:]:
    if r[0] and not any(c['id'] == r[0] for c in lab_caps):
        lab_caps.append({
            "id": r[0],
            "designationFamily": "LAB",
            "name": r[1],
            "type": r[2],
            "purpose": f"Advanced laboratory capability: {r[1]}",
            "category": "QUALITY"
        })

lab_features = []
for r in list(wb_lab1['Feature Catalog'].iter_rows(values_only=True))[1:]:
    if r[0]:
        lab_features.append({
            "designationFamily": "LAB",
            "featureType": r[0],
            "name": r[1],
            "requiredCapability": r[2],
            "accessLevel": r[3],
            "permissionId": r[4]
        })

for r in list(wb_lab2['Part 2 Feature Catalog'].iter_rows(values_only=True))[1:]:
    if r[0] and not any(f['permissionId'] == r[4] and f['name'] == r[1] for f in lab_features):
        lab_features.append({
            "designationFamily": "LAB",
            "featureType": r[0],
            "name": r[1],
            "requiredCapability": r[2],
            "accessLevel": r[3],
            "permissionId": r[4]
        })

lab_onboarding = []
for r in list(wb_lab1['Onboarding Mapping'].iter_rows(values_only=True))[1:]:
    if r[0]:
        lab_onboarding.append({
            "designationFamily": "LAB",
            "question": r[0],
            "answerSignal": r[1],
            "capabilityTrigger": r[2],
            "resultingFeatures": r[3]
        })

for r in list(wb_lab2['Part 2 Onboarding Mapping'].iter_rows(values_only=True))[1:]:
    if r[0]:
        lab_onboarding.append({
            "designationFamily": "LAB",
            "question": r[0],
            "answerSignal": r[1],
            "capabilityTrigger": r[2],
            "resultingFeatures": r[3]
        })

lab_policies = []
for r in list(wb_lab1['Lab Policy'].iter_rows(values_only=True))[1:]:
    if r[0]:
        lab_policies.append({
            "designationFamily": "LAB",
            "policy": r[0],
            "rule": r[1]
        })

lab_workspace_variants = []
for r in list(wb_lab2['Lab Workspace Variants'].iter_rows(values_only=True))[1:]:
    if r[0]:
        lab_workspace_variants.append({
            "designationFamily": "LAB",
            "variantId": r[0],
            "workspaceName": r[1],
            "operationalFocus": r[2],
            "capabilityCount": r[3],
            "capabilities": [c.strip() for c in r[4].split('+')],
            "navigation": [n.strip() for n in r[5].split('|')],
            "dashboardComposition": [d.strip() for d in r[6].split('|')]
        })

# 5. Extract Distributor Capabilities, Features, Variants, Policies
wb_dist = openpyxl.load_workbook('HoneyChain_Dispatch_Distributor_Feature_Matrix.xlsx', read_only=True)

dist_caps = []
for r in list(wb_dist['Capability Catalog'].iter_rows(values_only=True))[1:]:
    if r[0]:
        dist_caps.append({
            "id": r[0],
            "designationFamily": "DISTRIBUTOR",
            "name": r[1],
            "type": r[2],
            "purpose": r[3],
            "category": "FULFILLMENT"
        })

dist_features = []
for r in list(wb_dist['Feature Catalog'].iter_rows(values_only=True))[1:]:
    if r[0]:
        dist_features.append({
            "designationFamily": "DISTRIBUTOR",
            "featureType": r[0],
            "name": r[1],
            "requiredCapability": r[2],
            "accessLevel": r[3],
            "permissionId": r[4]
        })

dist_onboarding = []
for r in list(wb_dist['Onboarding Mapping'].iter_rows(values_only=True))[1:]:
    if r[0]:
        dist_onboarding.append({
            "designationFamily": "DISTRIBUTOR",
            "question": r[0],
            "answerSignal": r[1],
            "capabilityTrigger": r[2],
            "resultingFeatures": r[3]
        })

dist_policies = []
for r in list(wb_dist['Distribution Policy'].iter_rows(values_only=True))[1:]:
    if r[0]:
        dist_policies.append({
            "designationFamily": "DISTRIBUTOR",
            "policy": r[0],
            "rule": r[1]
        })

dist_workspace_variants = []
for r in list(wb_dist['Workspace Variants'].iter_rows(values_only=True))[1:]:
    if r[0]:
        dist_workspace_variants.append({
            "designationFamily": "DISTRIBUTOR",
            "variantId": r[0],
            "workspaceName": r[1],
            "operationalFocus": r[2],
            "capabilityCount": r[3],
            "capabilities": [c.strip() for c in r[4].split('+')],
            "navigation": [n.strip() for n in r[5].split('|')],
            "dashboardComposition": [d.strip() for d in r[6].split('|')]
        })

# Merge and Save Everything
all_capabilities = beekeeper_caps + processor_caps + lab_caps + dist_caps
all_features = beekeeper_features + processor_features + lab_features + dist_features
all_onboarding = beekeeper_onboarding + processor_onboarding + lab_onboarding + dist_onboarding
all_policies = beekeeper_policies + processor_policies + lab_policies + dist_policies
all_workspace_variants = lab_workspace_variants + dist_workspace_variants

# Add Beekeeper and Processor standard variants
all_workspace_variants.extend([
    {
        "designationFamily": "BEEKEEPER",
        "variantId": "BK-001",
        "workspaceName": "Beekeeper Basic",
        "operationalFocus": "Routine hive inspections and colony care",
        "capabilityCount": 4,
        "capabilities": ["HIVE_MANAGEMENT", "HIVE_INSPECTION", "BEE_OBSERVATION", "EVIDENCE_CAPTURE"],
        "navigation": ["Home", "Hives", "Inspections", "More"],
        "dashboardComposition": ["Colony conditions", "Recent inspections", "Quick actions"]
    },
    {
        "designationFamily": "BEEKEEPER",
        "variantId": "BK-002",
        "workspaceName": "Beekeeper Health Specialist",
        "operationalFocus": "Hive care with comb frame camera health screening",
        "capabilityCount": 5,
        "capabilities": ["HIVE_MANAGEMENT", "HIVE_INSPECTION", "BEE_OBSERVATION", "EVIDENCE_CAPTURE", "BEE_HEALTH_SCAN"],
        "navigation": ["Home", "Hives", "Inspections", "More"],
        "dashboardComposition": ["Colony conditions", "Frame health diagnostic", "Recent inspections", "Quick actions"]
    },
    {
        "designationFamily": "BEEKEEPER",
        "variantId": "BK-003",
        "workspaceName": "Connected Beekeeper",
        "operationalFocus": "Hive stewardship with ESP32 sensor telemetry",
        "capabilityCount": 6,
        "capabilities": ["HIVE_MANAGEMENT", "HIVE_INSPECTION", "BEE_OBSERVATION", "EVIDENCE_CAPTURE", "CONNECTED_HIVE_MONITORING", "HIVE_DEVICE_MANAGEMENT"],
        "navigation": ["Home", "Hives", "Inspections", "Devices", "More"],
        "dashboardComposition": ["Colony conditions", "Live telemetry", "Device status", "Quick actions"]
    },
    {
        "designationFamily": "BEEKEEPER",
        "variantId": "BK-004",
        "workspaceName": "Harvest Beekeeper",
        "operationalFocus": "Apiary care and honey supers harvesting",
        "capabilityCount": 6,
        "capabilities": ["HIVE_MANAGEMENT", "HIVE_INSPECTION", "BEE_OBSERVATION", "EVIDENCE_CAPTURE", "HONEY_COLLECTION", "COLLECTION_BATCH_LINK"],
        "navigation": ["Home", "Hives", "Inspections", "Honey", "More"],
        "dashboardComposition": ["Colony conditions", "Harvest collections", "Quick actions"]
    },
    {
        "designationFamily": "PROCESSOR",
        "variantId": "PR-001",
        "workspaceName": "Extraction Operator",
        "operationalFocus": "Batch intake, centrifugal extraction and settling tanks",
        "capabilityCount": 7,
        "capabilities": ["PROCESSING_MANAGEMENT", "BATCH_INTAKE", "PROCESSING_STEP_RECORD", "PROCESSING_PARAMETERS", "PROCESSING_EVIDENCE", "BATCH_TRACEABILITY", "PROCESSING_COMPLETION"],
        "navigation": ["Home", "Batches", "Processing", "More"],
        "dashboardComposition": ["Active curing tanks", "Batch intake queue", "Processing queue", "Quick actions", "Recent activity"]
    },
    {
        "designationFamily": "PROCESSOR",
        "variantId": "PR-002",
        "workspaceName": "Lead Processor",
        "operationalFocus": "Processing with batch split/merge and quality handoff",
        "capabilityCount": 10,
        "capabilities": ["PROCESSING_MANAGEMENT", "BATCH_INTAKE", "PROCESSING_STEP_RECORD", "PROCESSING_PARAMETERS", "PROCESSING_EVIDENCE", "BATCH_TRACEABILITY", "PROCESSING_COMPLETION", "BATCH_SPLIT_MERGE", "QUALITY_HANDOFF", "PACKAGING_HANDOFF"],
        "navigation": ["Home", "Batches", "Processing", "Handoffs", "More"],
        "dashboardComposition": ["Active curing tanks", "Processing queue", "Quality handoff status", "Packaging handoff status", "Quick actions", "Traceability"]
    }
])

with open('src/data/matrices/capabilities.json', 'w', encoding='utf-8') as f:
    json.dump(all_capabilities, f, indent=2)

with open('src/data/matrices/features.json', 'w', encoding='utf-8') as f:
    json.dump(all_features, f, indent=2)

with open('src/data/matrices/onboarding_tree.json', 'w', encoding='utf-8') as f:
    json.dump(all_onboarding, f, indent=2)

with open('src/data/matrices/policies.json', 'w', encoding='utf-8') as f:
    json.dump(all_policies, f, indent=2)

with open('src/data/matrices/workspace_variants.json', 'w', encoding='utf-8') as f:
    json.dump(all_workspace_variants, f, indent=2)

print(f"Extracted successfully: {len(all_capabilities)} capabilities, {len(all_features)} features/tools, {len(all_onboarding)} onboarding triggers, {len(all_workspace_variants)} workspace variants.")
